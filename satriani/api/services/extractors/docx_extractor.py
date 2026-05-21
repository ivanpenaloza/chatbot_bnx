"""Extractor for .docx files using python-docx.

Compatible with Python 3.9.20.
"""


class DocxExtractor:
    """Extract text from Word (.docx) documents."""

    def extract(self, filepath: str) -> str:
        from docx import Document

        doc = Document(filepath)
        paragraphs = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                paragraphs.append(text)

        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip() for cell in row.cells if cell.text.strip()
                )
                if row_text:
                    paragraphs.append(row_text)

        return "\n".join(paragraphs)
